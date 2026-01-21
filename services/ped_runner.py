from __future__ import annotations

import json
import re
import time
import unicodedata
import warnings
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import pandas as pd
from rapidfuzz import fuzz, process
from sqlalchemy import text
from sqlalchemy.exc import SQLAlchemyError

from models import db, Dotacao

# Evita warnings de downcasting silencioso em replace
pd.set_option("future.no_silent_downcasting", True)

BATCH_SIZE = 200

# Caminhos base
INPUT_DIR = Path("upload/ped")
OUTPUT_DIR = Path("outputs/td_ped")
JSON_CHAVES_PLANEJAMENTO = Path("static/js/chaves_planejamento.json")
JSON_FORCAR_CHAVE = Path("static/js/forcar_chave.json")
JSON_CASOS_ESPECIFICOS = Path("static/js/chave_arrumar.json")

# Cabeçalho mínimo esperado (normalizado)
HEADER_PADRAO_NORMALIZADO = ["EXERCICIO", "NO PED", "NO PED ESTORNO ESTORNADO", "NO EMP", "NO CAD", "NO NOBLIST", "NO OS"]
HEADER_PADRAO_PREFIXOS = ["EXERCICIO", "N", "N", "N", "N", "N", "N"]

# Nomes de colunas canônicos (lower normalizado -> nome final)
COLUNAS_CANONICAS = {
    "exercicio": "Exercício",
    "historico": "Histórico",
    "n ped": "Nº PED",
    "n ped estorno estornado": "Nº PED Estorno/Estornado",
    "n emp": "Nº EMP",
    "n cad": "Nº CAD",
    "n noblist": "Nº NOBLIST",
    "n os": "Nº OS",
    "dotacao orcamentaria": "Dotação Orçamentária",
    "data solicitacao": "Data Solicitação",
    "data criacao": "Data Criação",
    "data autorizacao": "Data Autorização",
    "data da licitacao": "Data da Licitação",
    "data hora cadastro autorizacao": "Data/Hora Cadastro Autorização",
    "exercicio de competencia da folha de pagamento": "Exercício de Competência da Folha de Pagamento",
    "natureza de despesa": "Natureza de Despesa",
}


def ensure_dirs() -> None:
    for base in (INPUT_DIR, OUTPUT_DIR, INPUT_DIR / "tmp", OUTPUT_DIR / "tmp"):
        base.mkdir(parents=True, exist_ok=True)


def move_existing_to_tmp(base_dir: Path) -> None:
    tmp = base_dir / "tmp"
    tmp.mkdir(parents=True, exist_ok=True)
    for f in base_dir.glob("*.xlsx"):
        dest = tmp / f"{f.stem}_{datetime.now().strftime('%Y%m%d%H%M%S')}{f.suffix}"
        try:
            f.rename(dest)
        except OSError:
            pass


def limpar_historico(texto: str) -> str:
    if not isinstance(texto, str):
        return "NÃO INFORMADO"
    texto = texto.replace("_x000D_", " ").replace("\n", " ").replace("\r", " ")
    texto = re.sub(r"\s+\*\s+", " * ", texto)
    texto = re.sub(r"\s+", " ", texto).strip()
    return texto if texto else "NÃO INFORMADO"


def corrigir_caracteres(texto: str) -> str:
    if not isinstance(texto, str):
        return "NÃO INFORMADO"
    texto = re.sub(r"[^\w\s,./\-|*]", "", texto)
    texto = re.sub(r"\s+", " ", texto).strip()
    return texto if texto else "NÃO INFORMADO"


def canonizar_nome_coluna(nome: str) -> str:
    if not isinstance(nome, str):
        return nome
    nome_norm = unicodedata.normalize("NFKD", nome)
    nome_norm = "".join(ch for ch in nome_norm if not unicodedata.combining(ch))
    nome_norm = re.sub(r"[^a-zA-Z0-9]+", " ", nome_norm).strip().lower()
    return COLUNAS_CANONICAS.get(nome_norm, nome.strip())


def extrair_ano(valor: Any) -> int | None:
    if valor is None:
        return None
    s = str(valor).strip()
    if not s:
        return None
    digits = re.sub(r"\D", "", s)
    if len(digits) >= 4:
        return int(digits[-4:])
    return None


def contar_partes_chave(chave: Any) -> int:
    if not isinstance(chave, str):
        return 0
    texto = chave.strip()
    if texto in ("", "-", "NÃO INFORMADO", "NÃO IDENTIFICADO", "NÇO INFORMADO", "NÇO IDENTIFICADO"):
        return 0
    if texto.upper().startswith("DOT."):
        base = texto.rstrip("*")
        partes = [p for p in base.split(".") if p.strip()]
        return len(partes)
    partes = [p.strip() for p in texto.split("*") if p.strip()]
    return len(partes)


def normalizar_colunas(df: pd.DataFrame) -> pd.DataFrame:
    return df.rename(columns={col: canonizar_nome_coluna(col) for col in df.columns})


def encontrar_coluna_prefixo(df: pd.DataFrame, prefixo: str) -> str | None:
    prefixo = (prefixo or "").lower()
    for col in df.columns:
        if isinstance(col, str) and col.lower().startswith(prefixo):
            return col
    return None


def _normalize_dotacao_key(value: str) -> str:
    if not value:
        return ""
    cleaned = re.sub(r"\s+", "", str(value)).rstrip("*")
    return cleaned.upper()


def _to_decimal(value: Any) -> Decimal:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return Decimal("0")
    raw = str(value).strip()
    if not raw:
        return Decimal("0")
    cleaned = raw.replace(".", "").replace(",", ".")
    try:
        return Decimal(cleaned)
    except Exception:
        return Decimal("0")


def _find_valor_ped_col(df: pd.DataFrame) -> str | None:
    for col in df.columns:
        if not isinstance(col, str):
            continue
        base = unicodedata.normalize("NFKD", col)
        base = "".join(ch for ch in base if not unicodedata.combining(ch))
        base = re.sub(r"[^A-Z0-9]+", "", base.upper())
        if base == "VALORPED":
            return col
    return None


def _update_dotacao_from_ped(df: pd.DataFrame) -> None:
    if "Chave" not in df.columns:
        return
    valor_col = _find_valor_ped_col(df)
    if not valor_col:
        return
    ped_sums: dict[str, Decimal] = {}
    for _, row in df.iterrows():
        chave = row.get("Chave")
        if not isinstance(chave, str):
            continue
        if not chave.strip().upper().startswith("DOT."):
            continue
        key = _normalize_dotacao_key(chave)
        ped_sums[key] = ped_sums.get(key, Decimal("0")) + _to_decimal(row.get(valor_col))

    dotacoes = Dotacao.query.filter(Dotacao.ativo == True).all()  # noqa: E712
    for dot in dotacoes:
        key = _normalize_dotacao_key(dot.chave_dotacao or "")
        ped_sum = ped_sums.get(key, Decimal("0"))
        dot_val = _to_decimal(dot.valor_dotacao)
        dot.valor_ped_emp = ped_sum
        dot.valor_atual = dot_val - ped_sum
    db.session.commit()


def carregar_chaves_planejamento(json_path: Path) -> list[str]:
    try:
        with open(json_path, "r", encoding="utf-8-sig") as file:
            chaves = json.load(file)
        return [corrigir_caracteres(re.sub(r"\s+", " ", chave.strip())) for chave in chaves]
    except Exception as e:
        print(f"Erro ao carregar as chaves de planejamento: {e}")
        return []


def carregar_casos_especificos(json_path: Path) -> dict[str, str]:
    try:
        with open(json_path, "r", encoding="utf-8-sig") as file:
            casos_especificos = json.load(file)
        casos_especificos = {corrigir_caracteres(k): corrigir_caracteres(v) for k, v in casos_especificos.items()}
        print(f"Casos específicos carregados do arquivo: {json_path}")
        return casos_especificos
    except Exception as e:
        print(f"Erro ao carregar casos específicos: {e}")
        return {}


def carregar_forcar_chave(json_path: Path) -> dict[str, str]:
    try:
        with open(json_path, "r", encoding="utf-8-sig") as file:
            mapping = json.load(file)
        return {str(k).strip(): corrigir_caracteres(str(v).strip()) for k, v in mapping.items()}
    except Exception as e:
        print(f"Erro ao carregar forcar_chave: {e}")
        return {}


def extrair_chave_valida_do_historico(hist_limpo: str, chaves_planejamento: list[str]) -> str | None:
    for chave in chaves_planejamento:
        if chave in hist_limpo:
            return chave
    return None


def identificar_chave_planejamento(df: pd.DataFrame, chaves_planejamento: list[str], casos_especificos: dict[str, str]) -> pd.DataFrame:
    def _vazio_emp_ou_estorno(v: Any) -> bool:
        if v is None or (isinstance(v, (int, float)) and v == 0):
            return True
        v = str(v).strip().upper()
        return v in (
            "",
            "NÃO INFORMADO",
            "NAO INFORMADO",
            "NÇO INFORMADO",
            "N€O INFORMADO",
            "N?O INFORMADO",
            "-",
            "0",
            "0.0",
            "0,0",
        )

    def encontrar_chave(row: pd.Series) -> str:
        hist = row.get("Histórico", "")
        ped_estorno = str(row.get("Nº PED Estorno/Estornado", "")).strip().upper()
        num_emp = str(row.get("Nº EMP", "")).strip().upper()
        exercicio_val = None
        for k in row.index:
            if isinstance(k, str) and k.lower().startswith("exerc"):
                exercicio_val = row.get(k)
                break
        ano = extrair_ano(exercicio_val)
        partes_planejamento = 8 if (ano and ano >= 2026) else 7

        if (not _vazio_emp_ou_estorno(ped_estorno)) or (not _vazio_emp_ou_estorno(num_emp)):
            return "IGNORADO"

        if hist == "NÃO INFORMADO":
            return "NÃO IDENTIFICADO"

        hist_text = str(hist or "")
        dot_match = re.search(r"\bDOT\.(\d{4})\.([A-Z0-9_-]+)\.(\d+)\*", hist_text, re.IGNORECASE)
        if dot_match:
            ano_dot, adj, id_dot = dot_match.groups()
            return f"DOT.{ano_dot}.{adj.upper()}.{id_dot}*"

        hist_limpo = re.sub(r"\s+", " ", hist).strip()
        if not hist_limpo.startswith("*"):
            hist_limpo = "* " + hist_limpo
        if not hist_limpo.endswith("*"):
            hist_limpo += " *"
        hist_limpo = re.sub(r"\s*\*\s*", " * ", hist_limpo)

        partes_hist = [p.strip() for p in hist_limpo.split("*") if p.strip()]
        for i in range(len(partes_hist) - 3):
            if partes_hist[i].upper() != "DOT":
                continue
            adj = partes_hist[i + 1]
            ano = partes_hist[i + 2]
            id_dot = partes_hist[i + 3]
            if re.fullmatch(r"\d{4}", ano) and re.fullmatch(r"\d+", id_dot):
                return f"DOT.{ano}.{adj.upper()}.{id_dot}*"

        chaves_preferidas = [c for c in chaves_planejamento if contar_partes_chave(c) == partes_planejamento]
        chave_direta = extrair_chave_valida_do_historico(hist_limpo, chaves_preferidas)
        if not chave_direta:
            chave_direta = extrair_chave_valida_do_historico(hist_limpo, chaves_planejamento)
        if chave_direta:
            return chave_direta

        for caso, chave in casos_especificos.items():
            if caso in hist_limpo:
                return chave

        partes = re.findall(r"\*([^*]+)", hist_limpo)
        if len(partes) >= partes_planejamento:
            trecho = " * ".join(partes[:partes_planejamento])
            base = chaves_preferidas or chaves_planejamento
            match = process.extractOne(trecho, base, scorer=fuzz.WRatio, score_cutoff=95)
            if match:
                print(f"Chave aproximada identificada por fuzzy: {match[0]}")
                return match[0]
        return "NÃO IDENTIFICADO"

    df["Chave"] = df.apply(encontrar_chave, axis=1)
    return df


def forcar_chaves_manualmente(df: pd.DataFrame, substituicoes: dict[str, str]) -> pd.DataFrame:
    if "Nº PED" in df.columns and substituicoes:
        if "_forcar_chave" not in df.columns:
            df["_forcar_chave"] = False
        df["Nº PED"] = df["Nº PED"].astype(str).str.strip()
        mask = df["Nº PED"].isin(substituicoes.keys())
        if mask.any():
            df.loc[mask, "Chave de Planejamento"] = df.loc[mask, "Nº PED"].map(substituicoes)
            df.loc[mask, "_forcar_chave"] = True
    return df


def converter_tipos(df: pd.DataFrame) -> pd.DataFrame:
    colunas_monetarias = ["Valor PED", "Valor do Estorno"]
    colunas_datas = ["Data da Licitação", "Data Solicitação", "Data Criação", "Data Autorização", "Data/Hora Cadastro Autorização"]
    colunas_numericas = ["Exercício de Competência da Folha de Pagamento"]

    df.replace({"": "NÃO INFORMADO", None: "NÃO INFORMADO"}, inplace=True)

    def parse_valor_monetario(v: Any) -> float:
        if pd.isna(v):
            return 0.00
        s = str(v).strip()
        if s in ("", "NÃO INFORMADO"):
            return 0.00
        s_num = re.sub(r"[^\d,.-]", "", s)
        if "," in s_num:
            s_num = s_num.replace(".", "").replace(",", ".")
        try:
            return round(float(s_num), 2)
        except ValueError:
            return 0.00

    def formatar_real_ptbr(valor_float: float) -> str:
        s = f"{valor_float:,.2f}"
        return s.replace(",", "X").replace(".", ",").replace("X", ".")

    for col in colunas_monetarias:
        if col in df.columns:
            df[col] = df[col].apply(parse_valor_monetario)
            df[col] = df[col].apply(formatar_real_ptbr)

    def formatar_data_br(x: Any) -> str:
        if pd.isna(x):
            return "00/00/0000"
        if isinstance(x, pd.Timestamp):
            if x.hour == 0 and x.minute == 0 and x.second == 0 and x.microsecond == 0:
                return x.strftime("%d/%m/%Y")
            return x.strftime("%d/%m/%Y %H:%M:%S")
        return "00/00/0000"

    for col in colunas_datas:
        if col in df.columns:
            with warnings.catch_warnings():
                warnings.filterwarnings("ignore", message="Could not infer format")
                serie_str = df[col].astype(str).str.strip()
                mask_iso = serie_str.str.match(r"\d{4}-\d{2}-\d{2}")
                parsed = pd.Series(index=serie_str.index, dtype="datetime64[ns]")
                if mask_iso.any():
                    parsed.loc[mask_iso] = pd.to_datetime(serie_str[mask_iso], errors="coerce", dayfirst=False)
                if (~mask_iso).any():
                    parsed.loc[~mask_iso] = pd.to_datetime(serie_str[~mask_iso], errors="coerce", dayfirst=True)
                df[col] = parsed
            df[col] = df[col].apply(formatar_data_br)

    for col in colunas_numericas:
        if col in df.columns:
            df[col] = df[col].fillna(0)

    return df


def adicionar_novas_colunas(df: pd.DataFrame) -> pd.DataFrame:
    novas_colunas_planejamento = ["Região", "Subfunção + UG", "ADJ", "Macropolítica", "Pilar", "Eixo", "Política_Decreto"]
    novas_colunas_orcamentarias = [
        "Função",
        "Subfunção",
        "Programa de Governo",
        "PAOE",
        "Natureza de Despesa",
        "Cat.Econ",
        "Grupo",
        "Modalidade",
        "Fonte",
        "Iduso",
        "Elemento",
        "Nome do Elemento",
    ]
    ex_col = encontrar_coluna_prefixo(df, "exerc")

    for col in novas_colunas_planejamento:
        if col not in df.columns and ex_col:
            df.insert(df.columns.get_loc(ex_col), col, "N?O INFORMADO")

    if "Dotação Orçamentária" in df.columns:
        posicao_insercao = df.columns.get_loc("Dotação Orçamentária") + 1
        for col in novas_colunas_orcamentarias:
            if col not in df.columns:
                df.insert(posicao_insercao, col, "NÃO INFORMADO")
                posicao_insercao += 1

    if all(col in df.columns for col in ["Elemento", "Nome do Elemento", "Modalidade", "Fonte"]):
        colunas = df.columns.tolist()
        colunas.remove("Elemento")
        colunas.remove("Nome do Elemento")
        idx = colunas.index("Modalidade")
        colunas.insert(idx + 1, "Elemento")
        colunas.insert(idx + 2, "Nome do Elemento")
        df = df[colunas]

    hist_col = encontrar_coluna_prefixo(df, "hist")
    if hist_col:
        df[hist_col] = df[hist_col].apply(limpar_historico)
        colunas = df.columns.tolist()
        colunas.remove("Credor")
        colunas.remove("Nome do Credor")
        idx = colunas.index("Histórico")
        colunas.insert(idx + 1, "Credor")
        colunas.insert(idx + 2, "Nome do Credor")
        df = df[colunas]

    return df


def preencher_novas_colunas(df: pd.DataFrame) -> pd.DataFrame:
    def extrair_valores(chave: str, partes: int = 7) -> list[str]:
        if not isinstance(chave, str) or chave.strip() in ["", "NÃO IDENTIFICADO", "NÃO INFORMADO", "-"]:
            return ["NÃO INFORMADO"] * partes
        pedacos = [p.strip() for p in chave.split("*") if p.strip()]
        if len(pedacos) < partes:
            return ["NÃO INFORMADO"] * partes
        return pedacos[:partes]

    valores_extraidos = df["Chave"].apply(lambda x: extrair_valores(x, partes=7))
    valores_extraidos = pd.DataFrame(
        valores_extraidos.tolist(),
        columns=["Região", "Subfunção + UG", "ADJ", "Macropolítica", "Pilar", "Eixo", "Política_Decreto"],
        index=df.index,
    )
    df.update(valores_extraidos)

    def extrair_dotacao(dot: str) -> list[str]:
        if not isinstance(dot, str) or dot.strip() == "":
            return ["NÃO INFORMADO"] * 7
        partes = dot.split(".")
        partes += [""] * (11 - len(partes))
        return [partes[2], partes[3], partes[4], partes[5], partes[7], partes[8], partes[9]]

    if "Dotação Orçamentária" in df.columns:
        valores_dotacao = df["Dotação Orçamentária"].apply(lambda x: extrair_dotacao(x))
        df_dot = pd.DataFrame(
            valores_dotacao.tolist(),
            columns=["Função", "Subfunção", "Programa de Governo", "PAOE", "Natureza de Despesa", "Fonte", "Iduso"],
            index=df.index,
        )
        df.update(df_dot)

    def extrair_natureza(n: str) -> list[str]:
        if not isinstance(n, str) or len(n) < 4:
            return ["NÃO INFORMADO"] * 3
        return [n[0], n[1], n[2:4]]

    if "Natureza de Despesa" in df.columns:
        natureza = df["Natureza de Despesa"].apply(lambda x: extrair_natureza(x))
        df_nat = pd.DataFrame(natureza.tolist(), columns=["Cat.Econ", "Grupo", "Modalidade"], index=df.index)
        df.update(df_nat)

    return df


def ajustar_largura_colunas(writer: pd.ExcelWriter, df: pd.DataFrame, sheet_name: str, largura_historico: int = 120) -> None:
    worksheet = writer.sheets[sheet_name]
    hist_col = encontrar_coluna_prefixo(df, "hist")
    if hist_col:
        col_index = df.columns.get_loc(hist_col)
        worksheet.set_column(col_index, col_index, largura_historico)
    for i, col in enumerate(df.columns):
        if hist_col and col == hist_col:
            continue
        if df.empty:
            max_length = len(str(col)) + 2
        else:
            max_val = df[col].astype(str).map(len).max()
            if pd.isna(max_val):
                max_val = len(str(col))
            max_length = int(max(max_val, len(str(col))) + 2)
        worksheet.set_column(i, i, max_length)


def encontrar_linha_cabecalho(df_raw: pd.DataFrame) -> tuple[int, int] | None:
    def normalizar(texto: Any) -> str:
        if not isinstance(texto, str):
            return ""
        texto = unicodedata.normalize("NFKD", texto)
        texto = "".join(ch for ch in texto if not unicodedata.combining(ch))
        texto = re.sub(r"[^\w\s]", " ", texto)
        texto = re.sub(r"\s+", " ", texto).strip().upper()
        return texto

    def classificar(valor_norm: str) -> str | None:
        if not valor_norm:
            return None
        if "EXERCICIO" in valor_norm:
            return "exercicio"
        if "PED" in valor_norm and "EST" in valor_norm:
            return "ped_estorno"
        if "PED" in valor_norm:
            return "ped"
        if "EMP" in valor_norm:
            return "emp"
        if "CAD" in valor_norm:
            return "cad"
        if "NOBLIS" in valor_norm:
            return "noblist"
        if re.search(r"\bOS\b", valor_norm) or valor_norm.endswith(" OS") or valor_norm.startswith("N OS"):
            return "os"
        return None

    required = ["exercicio", "ped", "ped_estorno", "emp", "cad", "noblist"]
    emp_only = {"exercicio", "emp", "ped"}
    emp_only_headers = False

    max_rows = min(20, len(df_raw))
    for idx in range(max_rows):
        row = df_raw.iloc[idx]
        valores = [normalizar(row[i]) if pd.notna(row[i]) else "" for i in range(len(row))]
        if any("EXERCICIO IGUAL A" in v for v in valores):
            continue
        tipos = [classificar(v) for v in valores]
        tipos_set = {t for t in tipos if t}
        if emp_only.issubset(tipos_set):
            emp_only_headers = True
        posicoes: dict[str, int] = {}
        last_idx = -1
        for req in required:
            try:
                next_idx = next(i for i in range(last_idx + 1, len(tipos)) if tipos[i] == req)
            except StopIteration:
                break
            posicoes[req] = next_idx
            last_idx = next_idx
        if len(posicoes) == len(required):
            return idx, posicoes["exercicio"]

    # Fallback: aceita a linha se todos os campos obrigatorios aparecem em qualquer ordem.
    for idx in range(max_rows):
        row = df_raw.iloc[idx]
        valores = [normalizar(row[i]) if pd.notna(row[i]) else "" for i in range(len(row))]
        if any("EXERCICIO IGUAL A" in v for v in valores):
            continue
        tipos = [classificar(v) for v in valores]
        posicoes = {req: next((i for i, t in enumerate(tipos) if t == req), None) for req in required}
        if all(p is not None for p in posicoes.values()):
            return idx, min(posicoes.values())

    if emp_only_headers:
        raise RuntimeError("Arquivo não parece ser PED (cabeçalho de EMP detectado).")

    print("DEBUG: cabecalho esperado (normalizado):", HEADER_PADRAO_NORMALIZADO)
    limite = min(10, len(df_raw))
    for idx in range(limite):
        raw = [df_raw.iloc[idx, j] if j < df_raw.shape[1] else "" for j in range(len(HEADER_PADRAO_NORMALIZADO))]
        norm = [normalizar(val) for val in raw]
        print(f"DEBUG linha {idx}: raw={raw} | norm={norm}")

    return None


def preparar_aba_ped(file_path: Path) -> pd.DataFrame | None:
    try:
        xls = pd.ExcelFile(file_path)
        sheet_names = list(xls.sheet_names)
        try:
            from openpyxl import load_workbook

            wb = load_workbook(file_path, read_only=True, data_only=True)
            active_title = wb.active.title if wb.active else None
            if active_title in sheet_names:
                sheet_names.remove(active_title)
                sheet_names.insert(0, active_title)
        except Exception:
            pass

        if active_title:
            df_raw = pd.read_excel(xls, sheet_name=active_title, header=None, dtype=str)
            try:
                resultado = encontrar_linha_cabecalho(df_raw)
            except RuntimeError:
                raise
            if resultado is not None:
                idx_cabecalho, col_inicio = resultado
                cabecalho = [
                    str(c).strip() if pd.notna(c) else "" for c in df_raw.iloc[idx_cabecalho, col_inicio:].tolist()
                ]
                last_non_empty = 0
                for i in range(len(cabecalho) - 1, -1, -1):
                    if cabecalho[i]:
                        last_non_empty = i
                        break
                cabecalho = cabecalho[: last_non_empty + 1]
                df = df_raw.iloc[idx_cabecalho + 1 :, col_inicio : col_inicio + len(cabecalho)].copy()
                df.columns = cabecalho
                df = df.dropna(how="all")
                df = normalizar_colunas(df)
                return df

        for sheet_name in sheet_names:
            df_raw = pd.read_excel(xls, sheet_name=sheet_name, header=None, dtype=str)
            try:
                resultado = encontrar_linha_cabecalho(df_raw)
            except RuntimeError:
                raise
            if resultado is None:
                continue
            idx_cabecalho, col_inicio = resultado
            cabecalho = [
                str(c).strip() if pd.notna(c) else "" for c in df_raw.iloc[idx_cabecalho, col_inicio:].tolist()
            ]
            last_non_empty = 0
            for i in range(len(cabecalho) - 1, -1, -1):
                if cabecalho[i]:
                    last_non_empty = i
                    break
            cabecalho = cabecalho[: last_non_empty + 1]
            df = df_raw.iloc[idx_cabecalho + 1 :, col_inicio : col_inicio + len(cabecalho)].copy()
            df.columns = cabecalho
            df = df.dropna(how="all")
            df = normalizar_colunas(df)
            return df

        print(f"Cabecalho padrao nao encontrado em nenhuma aba de {file_path}")
        return None
    except Exception as e:
        print(f"Erro ao preparar aba 'ped' para {file_path}: {e}")
        return None


def prefiltrar_ped(df: pd.DataFrame) -> pd.DataFrame:
    def norm_col_name(col_name: str) -> str:
        nome = unicodedata.normalize("NFKD", col_name or "")
        nome = "".join(ch for ch in nome if not unicodedata.combining(ch))
        nome = re.sub(r"[^A-Z0-9]+", " ", nome.upper()).strip()
        return nome

    def _match_emp_col(name_norm: str) -> bool:
        tokens = name_norm.split()
        if not tokens:
            return False
        if tokens[0] not in ("N", "NO", "NUM", "NUMERO", "NRO"):
            return False
        return any(t.startswith("EMP") for t in tokens)

    def _match_estorno_col(name_norm: str) -> bool:
        tokens = name_norm.split()
        if not tokens:
            return False
        if tokens[0] not in ("N", "NO", "NUM", "NUMERO", "NRO"):
            return False
        return ("PED" in tokens or any(t.startswith("PED") for t in tokens)) and any(t.startswith("ESTORN") for t in tokens)

    colunas_norm = {c: norm_col_name(c) for c in df.columns if isinstance(c, str)}
    estorno_col = next((c for c, n in colunas_norm.items() if n == "N PED ESTORNO ESTORNADO"), None)
    emp_col = next((c for c, n in colunas_norm.items() if n == "N EMP"), None)
    if not estorno_col:
        estorno_col = next((c for c, n in colunas_norm.items() if _match_estorno_col(n)), None)
    if not emp_col:
        emp_col = next((c for c, n in colunas_norm.items() if _match_emp_col(n)), None)

    def _is_vazio_ou_zero_raw(valor: Any, aceita_hifen: bool) -> bool:
        if pd.isna(valor):
            return True
        if isinstance(valor, (int, float)) and valor == 0:
            return True
        s = str(valor).strip()
        if s == "":
            return True
        upper = s.upper()
        if upper in ("NAN", "NONE", "NÃO INFORMADO", "NAO INFORMADO", "NÇO INFORMADO"):
            return True
        if aceita_hifen and s == "-":
            return True
        if re.fullmatch(r"-?\d+(?:[.,]\d+)?", s):
            try:
                return float(s.replace(",", ".")) == 0.0
            except ValueError:
                return False
        return False

    if estorno_col:
        df = df[df[estorno_col].apply(lambda v: _is_vazio_ou_zero_raw(v, aceita_hifen=True))]

    if emp_col:
        df = df[df[emp_col].apply(lambda v: _is_vazio_ou_zero_raw(v, aceita_hifen=True))]

    return df


def processar_planilha(
    df: pd.DataFrame, chaves_planejamento: list[str], casos_especificos: dict[str, str], forcar_map: dict[str, str]
) -> pd.DataFrame | None:
    try:
        ano = None
        ex_col = encontrar_coluna_prefixo(df, "exerc")
        if ex_col:
            anos = df[ex_col].apply(extrair_ano).dropna()
            if not anos.empty:
                ano = int(anos.mode().iloc[0])

        df = prefiltrar_ped(df)

        hist_col = encontrar_coluna_prefixo(df, "hist")
        if hist_col:
            df[hist_col] = df[hist_col].apply(limpar_historico)
        cols_obj = df.select_dtypes(include=["object"]).columns
        df[cols_obj] = df[cols_obj].apply(lambda col: col.map(corrigir_caracteres))

        df = converter_tipos(df)
        df = identificar_chave_planejamento(df, chaves_planejamento, casos_especificos)

        if "Chave" in df.columns:
            colunas = df.columns.tolist()
            colunas.insert(0, colunas.pop(colunas.index("Chave")))
            df = df[colunas]

        partes_planejamento = 7
        if ano and ano >= 2026:
            partes_planejamento = 8

        precisa_colunas_planejamento = False
        if "Chave" in df.columns:
            partes = df["Chave"].apply(contar_partes_chave)
            precisa_colunas_planejamento = (partes >= 7).any()
        if precisa_colunas_planejamento:
            df = adicionar_novas_colunas(df)
        df = preencher_novas_colunas(df)

        # Ajusta colunas "Chave" vs "Chave de Planejamento" conforme ano e formato da chave
        def ajustar_chave_por_formato(row: pd.Series) -> pd.Series:
            if row.get("_forcar_chave"):
                return row
            chave = row.get("Chave", "")
            partes = contar_partes_chave(chave)
            if partes == partes_planejamento:
                row["Chave de Planejamento"] = chave or "-"
                row["Chave"] = "-"
            elif partes == 4:
                row["Chave"] = chave or "-"
                row["Chave de Planejamento"] = "-"
            else:
                row["Chave de Planejamento"] = row.get("Chave de Planejamento") or "-"
                row["Chave"] = row.get("Chave") or "-"
            return row

        df = df.apply(ajustar_chave_por_formato, axis=1)
        df = forcar_chaves_manualmente(df, forcar_map)

        df = df.replace(
            {
                "NÃO INFORMADO": "-",
                "NÃO IDENTIFICADO": "-",
                "NÇŸO INFORMADO": "-",
                "NÇŸO IDENTIFICADO": "-",
                "NÇO INFORMADO": "-",
                "NÇO IDENTIFICADO": "-",
                "N€YO INFORMADO": "-",
                "N€YO IDENTIFICADO": "-",
            },
            regex=False,
        )

        return df
    except Exception as e:
        print(f"Erro ao processar a planilha: {e}")
        return None


def salvar_planilhas(ped_df: pd.DataFrame, tratado_df: pd.DataFrame, file_path: Path) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    move_existing_to_tmp(OUTPUT_DIR)

    output_file = OUTPUT_DIR / f"{file_path.stem}_Tratado.xlsx"
    writer = pd.ExcelWriter(output_file, engine="xlsxwriter")

    ped_df.to_excel(writer, index=False, sheet_name="ped")
    tratado_df.to_excel(writer, index=False, sheet_name="ped_tratado")

    ajustar_largura_colunas(writer, ped_df, "ped", largura_historico=60)
    ajustar_largura_colunas(writer, tratado_df, "ped_tratado", largura_historico=120)

    writer.close()
    return output_file


DF_TO_DB = {
    "Chave": "chave",
    "Chave de Planejamento": "chave_planejamento",
    "Região": "regiao",
    "Subfunção + UG": "subfuncao_ug",
    "ADJ": "adj",
    "Macropolítica": "macropolitica",
    "Pilar": "pilar",
    "Eixo": "eixo",
    "Política_Decreto": "politica_decreto",
    "Exercício": "exercicio",
    "Histórico": "historico",
    "Nº PED": "numero_ped",
    "Nº PED Estorno/Estornado": "numero_ped_estorno",
    "Nº EMP": "numero_emp",
    "Nº CAD": "numero_cad",
    "Nº NOBLIST": "numero_noblist",
    "Nº OS": "numero_os",
    "Convênio": "convenio",
    "Nº Processo Orçamentário de Pagamento": "numero_processo_orcamentario_pagamento",
    "Valor PED": "valor_ped",
    "Valor do Estorno": "valor_estorno",
    "Indicativo de Licitação de Exercícios Anteriores": "indicativo_licitacao_exercicios_anteriores",
    "Data da Licitação": "data_licitacao",
    "Liberado Fisco Estadual": "liberado_fisco_estadual",
    "Situação": "situacao",
    "UO": "uo",
    "Nome da Unidade Orçamentária": "nome_unidade_orcamentaria",
    "UG": "ug",
    "Nome da Unidade Gestora": "nome_unidade_gestora",
    "Data Solicitação": "data_solicitacao",
    "Data Criação": "data_criacao",
    "Tipo Empenho": "tipo_empenho",
    "Dotação Orçamentária": "dotacao_orcamentaria",
    "Função": "funcao",
    "Subfunção": "subfuncao",
    "Programa de Governo": "programa_governo",
    "PAOE": "paoe",
    "Natureza de Despesa": "natureza_despesa",
    "Cat.Econ": "cat_econ",
    "Grupo": "grupo",
    "Modalidade": "modalidade",
    "Elemento": "elemento",
    "Nome do Elemento": "nome_elemento",
    "Fonte": "fonte",
    "Iduso": "iduso",
    "Nº Emenda (EP)": "numero_emenda_ep",
    "Autor da Emenda (EP)": "autor_emenda_ep",
    "Nº CAC": "numero_cac",
    "Licitação": "licitacao",
    "Usuário Responsável": "usuario_responsavel",
    "Credor": "credor",
    "Nome do Credor": "nome_credor",
    "Data Autorização": "data_autorizacao",
    "Data/Hora Cadastro Autorização": "data_hora_cadastro_autorizacao",
    "Tipo de Despesa": "tipo_despesa",
    "Nº ABJ": "numero_abj",
    "Nº Processo do Sequestro Judicial": "numero_processo_sequestro_judicial",
    "Indicativo de Entrega imediata - § 4º  Art. 62 Lei 8.666": "indicativo_entrega_imediata",
    "Indicativo de contrato": "indicativo_contrato",
    "Código UO Extinta": "codigo_uo_extinta",
    "Devolução GCV": "devolucao_gcv",
    "Mês de Competência da Folha de Pagamento": "mes_competencia_folha_pagamento",
    "Exercício de Competência da Folha de Pagamento": "exercicio_competencia_folha",
    "Obrigação Patronal": "obrigacao_patronal",
    "Tipo de Obrigação Patronal": "tipo_obrigacao_patronal",
    "Nº NLA": "numero_nla",
}


def _clean_val(val: Any) -> Any:
    try:
        if pd.isna(val):
            return None
    except Exception:
        pass
    if isinstance(val, str) and val.strip() == "-":
        return None
    return val


def _parse_valor_db(valor: Any) -> float | None:
    """
    Converte strings formatadas (pt-BR) em float para inserir em colunas numéricas.
    Retorna None para valores vazios/inválidos.
    """
    if valor is None:
        return None
    s = str(valor).strip()
    if s in ("", "-", "NÃO INFORMADO", "NÃO IDENTIFICADO"):
        return None
    s_num = re.sub(r"[^\d,.-]", "", s)
    if "," in s_num:
        s_num = s_num.replace(".", "").replace(",", ".")
    try:
        return float(s_num)
    except ValueError:
        return None


def _parse_data_db(valor: Any) -> datetime | None:
    """
    Converte datas em dd/mm/yyyy (opcional hh:mm:ss) para datetime ou retorna None.
    Evita inserir placeholders como 00/00/0000.
    """
    if valor is None:
        return None
    if isinstance(valor, datetime):
        return valor
    s = str(valor).strip()
    if not s or s in ("-", "00/00/0000", "00/00/0000 00:00:00"):
        return None
    # normaliza separador
    s = s.replace("-", "/")
    for fmt in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def montar_registros_para_db(df: pd.DataFrame, data_arquivo: datetime, user_email: str, upload_id: int) -> list[dict[str, Any]]:
    registros: list[dict[str, Any]] = []
    rows = df.to_dict(orient="records")
    for row in rows:
        exercicio_val = None
        for k in row:
            if isinstance(k, str) and k.lower().startswith("exerc"):
                exercicio_val = row.get(k)
                break
        ano = extrair_ano(exercicio_val)
        payload: dict[str, Any] = {}
        for col_df, col_db in DF_TO_DB.items():
            payload[col_db] = _clean_val(row.get(col_df))
        if ano:
            payload["exercicio"] = str(ano)
        # Ajuste de chave x chave_planejamento conforme ano e formato da chave
        chave_val = row.get("Chave")
        chave_planejamento_val = row.get("Chave de Planejamento")
        if row.get("_forcar_chave"):
            payload["chave"] = _clean_val(chave_val)
            payload["chave_planejamento"] = _clean_val(chave_planejamento_val)
        else:
            partes = contar_partes_chave(chave_val)
            partes_planejamento = 7
            if ano and ano >= 2026:
                partes_planejamento = 8
            if partes == partes_planejamento:
                payload["chave_planejamento"] = _clean_val(chave_val)
                payload["chave"] = None
            elif partes == 4:
                payload["chave"] = _clean_val(chave_val)
                payload["chave_planejamento"] = None
            else:
                payload["chave"] = _clean_val(chave_val)
                payload["chave_planejamento"] = _clean_val(chave_planejamento_val)
        # Campos monetarios em float para evitar erro de conversao no DB
        if "valor_ped" in payload:
            payload["valor_ped"] = _parse_valor_db(payload["valor_ped"])
        if "valor_estorno" in payload:
            payload["valor_estorno"] = _parse_valor_db(payload["valor_estorno"])
        # Campos de data convertidos para datetime ou None
        for k in (
            "data_solicitacao",
            "data_criacao",
            "data_autorizacao",
            "data_licitacao",
            "data_hora_cadastro_autorizacao",
        ):
            if k in payload:
                payload[k] = _parse_data_db(payload[k])
        payload["upload_id"] = upload_id
        payload["data_atualizacao"] = datetime.utcnow()
        payload["data_arquivo"] = data_arquivo
        payload["user_email"] = user_email
        payload["ativo"] = True
        registros.append(payload)
    return registros

def update_database(df: pd.DataFrame, data_arquivo: datetime, user_email: str, upload_id: int) -> int:
    insert_sql = text(
        """
        INSERT INTO ped (
            upload_id, chave, regiao, subfuncao_ug, adj, macropolitica, pilar, eixo, politica_decreto,
            exercicio, historico, numero_ped, numero_ped_estorno, numero_emp, numero_cad, numero_noblist,
            numero_os, convenio, indicativo_licitacao_exercicios_anteriores, liberado_fisco_estadual, situacao,
            uo, nome_unidade_orcamentaria, ug, nome_unidade_gestora, numero_processo_orcamentario_pagamento,
            valor_ped, valor_estorno, dotacao_orcamentaria, funcao, subfuncao, programa_governo, paoe,
            natureza_despesa, cat_econ, grupo, modalidade, elemento, nome_elemento, fonte, iduso,
            numero_emenda_ep, autor_emenda_ep, numero_cac, licitacao, usuario_responsavel, data_solicitacao,
            data_criacao, data_autorizacao, data_licitacao, data_hora_cadastro_autorizacao, tipo_empenho,
            tipo_despesa, numero_abj, numero_processo_sequestro_judicial, indicativo_entrega_imediata,
            indicativo_contrato, codigo_uo_extinta, devolucao_gcv, mes_competencia_folha_pagamento,
            exercicio_competencia_folha, obrigacao_patronal, tipo_obrigacao_patronal, numero_nla, credor,
            nome_credor, chave_planejamento, data_atualizacao, data_arquivo, user_email, ativo
        )
        VALUES (
            :upload_id, :chave, :regiao, :subfuncao_ug, :adj, :macropolitica, :pilar, :eixo, :politica_decreto,
            :exercicio, :historico, :numero_ped, :numero_ped_estorno, :numero_emp, :numero_cad, :numero_noblist,
            :numero_os, :convenio, :indicativo_licitacao_exercicios_anteriores, :liberado_fisco_estadual, :situacao,
            :uo, :nome_unidade_orcamentaria, :ug, :nome_unidade_gestora, :numero_processo_orcamentario_pagamento,
            :valor_ped, :valor_estorno, :dotacao_orcamentaria, :funcao, :subfuncao, :programa_governo, :paoe,
            :natureza_despesa, :cat_econ, :grupo, :modalidade, :elemento, :nome_elemento, :fonte, :iduso,
            :numero_emenda_ep, :autor_emenda_ep, :numero_cac, :licitacao, :usuario_responsavel, :data_solicitacao,
            :data_criacao, :data_autorizacao, :data_licitacao, :data_hora_cadastro_autorizacao, :tipo_empenho,
            :tipo_despesa, :numero_abj, :numero_processo_sequestro_judicial, :indicativo_entrega_imediata,
            :indicativo_contrato, :codigo_uo_extinta, :devolucao_gcv, :mes_competencia_folha_pagamento,
            :exercicio_competencia_folha, :obrigacao_patronal, :tipo_obrigacao_patronal, :numero_nla, :credor,
            :nome_credor, :chave_planejamento, :data_atualizacao, :data_arquivo, :user_email, :ativo
        )
        """
    )

    try:
        db.session.execute(text("UPDATE ped SET ativo = 0 WHERE ativo = 1"))
        db.session.commit()
    except SQLAlchemyError:
        db.session.rollback()
        raise

    registros = montar_registros_para_db(df, data_arquivo, user_email, upload_id)
    total = 0
    for start in range(0, len(registros), BATCH_SIZE):
        chunk = registros[start : start + BATCH_SIZE]
        try:
            db.session.execute(insert_sql, chunk)
            db.session.commit()
            total += len(chunk)
        except SQLAlchemyError:
            db.session.rollback()
            raise
    return total


def _normalize_dotacao_key(value: str) -> str:
    if not value:
        return ""
    cleaned = re.sub(r"\s+", "", str(value)).rstrip("*")
    return cleaned.upper()


def _find_missing_dotacao_keys(df: pd.DataFrame) -> list[str]:
    if "Chave" not in df.columns:
        return []
    dot_keys = {
        _normalize_dotacao_key(val)
        for val in df["Chave"]
        if isinstance(val, str) and val.strip().upper().startswith("DOT.")
    }
    dot_keys = {k for k in dot_keys if k}
    if not dot_keys:
        return []
    db_keys = (
        db.session.query(Dotacao.chave_dotacao)
        .filter(Dotacao.chave_dotacao.isnot(None))
        .all()
    )
    db_norm = {_normalize_dotacao_key(k[0]) for k in db_keys if k and k[0]}
    return sorted([k for k in dot_keys if k not in db_norm])


def run_ped(
    file_path: Path, data_arquivo: datetime, user_email: str, upload_id: int
) -> tuple[int, Path, list[str]]:
    ensure_dirs()
    chaves_planejamento = carregar_chaves_planejamento(JSON_CHAVES_PLANEJAMENTO)
    casos_especificos = carregar_casos_especificos(JSON_CASOS_ESPECIFICOS)
    forcar_map = carregar_forcar_chave(JSON_FORCAR_CHAVE)

    ped_df = preparar_aba_ped(file_path)
    if ped_df is None:
        raise RuntimeError("Falha ao identificar cabeçalho ou ler a aba ped.")

    tratado_df = processar_planilha(ped_df.copy(), chaves_planejamento, casos_especificos, forcar_map)
    if tratado_df is None:
        raise RuntimeError("Falha ao tratar a planilha PED.")

    missing_dotacao_keys = _find_missing_dotacao_keys(tratado_df)
    tratado_df_export = tratado_df.drop(columns=["_forcar_chave"], errors="ignore")
    output_path = salvar_planilhas(ped_df, tratado_df_export, file_path)
    try:
        total = update_database(tratado_df, data_arquivo, user_email, upload_id)
    except SQLAlchemyError as exc:
        if "Packet sequence number wrong" in str(exc):
            db.session.remove()
            try:
                db.engine.dispose()
            except Exception:
                pass
            total = update_database(tratado_df, data_arquivo, user_email, upload_id)
        else:
            raise
    _update_dotacao_from_ped(tratado_df)
    return total, output_path, missing_dotacao_keys
